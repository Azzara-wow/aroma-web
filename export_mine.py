# export_mine.py — aroma_web / выгрузка вкладки «Моё» в Excel или PDF.
#
# На входе «Моё» в одной форме (archive.mine_current / archive.mine_of):
#   {name, items: [{aroma, volume, amount, category, unit, problem?}], total, bill{...}}
# плюс наличие текущей закупки (если есть). PDF — reportlab со шрифтом с кириллицей.

import io
import os
from datetime import datetime

FONT_PATHS = [
    os.environ.get("PDF_FONT", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",   # сервер (Ubuntu)
    "C:/Windows/Fonts/arial.ttf",                         # локальная проверка
]
FONT_BOLD_PATHS = [
    os.environ.get("PDF_FONT_BOLD", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
]


def _rows(mine, nal_items):
    """Строки таблицы и итоги — общие для Excel и PDF."""
    rows = []
    for i, it in enumerate(mine.get("items", []) if mine else [], start=1):
        unit = it.get("unit") or "мл"
        rows.append([i, it["aroma"], it.get("category", ""), f'{it["volume"]} {unit}',
                     it["amount"] if not it.get("problem") else "уточняется"])
    total = (mine or {}).get("total", 0)
    lines = [("Заказ по закупке", total)]
    bill = (mine or {}).get("bill") or {}
    deliv = _int(bill.get("pay_delivery"))
    if deliv:
        lines.append(("Доставка", deliv))
    if bill.get("pay_amount"):
        lines.append(("Оплачено" if bill.get("paid") else "Итого к оплате", _int(bill["pay_amount"])))
    nal = []
    for it in nal_items or []:
        s = round(it["mine"] * it["per_ml"]) if it.get("per_ml") else ""
        nal.append([it["name"], it["mine"], s])
    return rows, lines, nal


def _int(v):
    try:
        return int(float(str(v).replace(",", ".").replace(" ", "")))
    except (TypeError, ValueError):
        return 0


def _stamp():
    return datetime.now().strftime("%d.%m.%Y %H:%M")


def xlsx(title, who, phone, mine, nal_items=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows, lines, nal = _rows(mine, nal_items)
    wb = Workbook()
    ws = wb.active
    ws.title = "Моё"
    ws.append([f"LUZI · {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{who} · {phone}"])
    ws.append([f"Выгружено {_stamp()}"])
    ws["A3"].font = Font(color="888888", size=9)
    ws.append([])
    head = ["№", "Аромат", "Категория", "Объём", "Сумма, ₽"]
    ws.append(head)
    thin = Side(style="thin", color="CCCCCC")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EFE4")
        c.border = Border(bottom=thin)
    for r in rows:
        ws.append(r)
    ws.append([])
    for label, v in lines:
        ws.append(["", label, "", "", v])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=5).font = Font(bold=True)
    if nal:
        ws.append([])
        ws.append(["", "Наличие (склад) — считается отдельно"])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        for name, qty, s in nal:
            ws.append(["", name, "", qty, s])
    for col, w in zip("ABCDE", (5, 42, 14, 12, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=6):
        row[4].alignment = Alignment(horizontal="right")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "LuziSans" in pdfmetrics.getRegisteredFontNames():
        return "LuziSans", "LuziSans-Bold"
    reg = next((p for p in FONT_PATHS if p and os.path.exists(p)), None)
    bold = next((p for p in FONT_BOLD_PATHS if p and os.path.exists(p)), reg)
    if not reg:
        raise RuntimeError("Нет шрифта с кириллицей для PDF (DejaVuSans.ttf)")
    pdfmetrics.registerFont(TTFont("LuziSans", reg))
    pdfmetrics.registerFont(TTFont("LuziSans-Bold", bold))
    return "LuziSans", "LuziSans-Bold"


def pdf(title, who, phone, mine, nal_items=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold = _font()
    green = colors.HexColor("#2c4f2c")
    st_h = ParagraphStyle("h", fontName=bold, fontSize=15, leading=19, textColor=green)
    st = ParagraphStyle("p", fontName=reg, fontSize=10, leading=13)
    st_s = ParagraphStyle("s", fontName=reg, fontSize=8, leading=10, textColor=colors.grey)
    cell = ParagraphStyle("c", fontName=reg, fontSize=9.5, leading=12)

    rows, lines, nal = _rows(mine, nal_items)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"LUZI · {title} · {who}")
    story = [Paragraph(f"LUZI · {title}", st_h), Spacer(1, 3),
             Paragraph(f"{who} · {phone}", st), Paragraph(f"Выгружено {_stamp()}", st_s), Spacer(1, 10)]

    data = [["№", "Аромат", "Категория", "Объём", "Сумма, ₽"]]
    data += [[r[0], Paragraph(str(r[1]), cell), r[2], r[3], r[4]] for r in rows]
    for label, v in lines:
        data.append(["", label, "", "", v])
    n_items = len(rows)
    t = Table(data, colWidths=[9 * mm, 82 * mm, 28 * mm, 22 * mm, 25 * mm], repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, -1), reg, 9.5),
        ("FONT", (0, 0), (-1, 0), bold, 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8efe4")),
        ("LINEBELOW", (0, 0), (-1, n_items), 0.3, colors.HexColor("#cccccc")),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if lines:
        style += [("FONT", (1, n_items + 1), (-1, -1), bold, 10),
                  ("LINEABOVE", (0, n_items + 1), (-1, n_items + 1), 0.8, green)]
    t.setStyle(TableStyle(style))
    story.append(t)

    if nal:
        story += [Spacer(1, 12), Paragraph("Наличие (склад) — считается отдельно", ParagraphStyle(
            "h2", fontName=bold, fontSize=11, leading=14))]
        nt = Table([["Товар", "Кол-во", "Сумма, ₽"]] + [[Paragraph(str(n), cell), q, s] for n, q, s in nal],
                   colWidths=[119 * mm, 22 * mm, 25 * mm])
        nt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), reg, 9.5), ("FONT", (0, 0), (-1, 0), bold, 9.5),
                                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc"))]))
        story.append(nt)
    if not rows and not nal:
        story.append(Paragraph("Заказов нет.", st))
    doc.build(story)
    return buf.getvalue()
